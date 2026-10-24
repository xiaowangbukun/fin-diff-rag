# SPDX-License-Identifier: Apache-2.0
"""文档接入：解析 → 语义切片 → 敏感信息统计 → 向量化 → 落库。"""
import hashlib
import re
from datetime import datetime
from pathlib import Path

from sqlalchemy import delete as sa_delete, select
from sqlalchemy.orm import Session

from .config import CHUNK_OVERLAP, CHUNK_SIZE, DATA_DIR, MAX_DOC_CHARS
from .embed import embedder
from .injection import scan_injection
from .models import Chunk, Document, DocumentVersion
from .pdfx import ParseError, assess_quality, parse_pdf
from .pii import scan

SENT_SPLIT = re.compile(r"(?<=[。！？!?；;\n])")

# 段落级密级标记：命中即把该切片密级上调（同文档跨密级隔离，治段落级撕裂）。
# PII 不在此列——那由脱敏链路负责，密级与脱敏是两个维度的控制。
CHUNK_MARKERS: tuple[tuple[str, int], ...] = (
    ("绝密", 3), ("机密", 2), ("保密", 2), ("内部限阅", 2), ("限制传阅", 2),
)


def detect_chunk_clearance(piece: str, doc_clearance: int) -> int | None:
    """探测切片级密级：显式密级标记 → 提升；无标记 → None（继承文档密级）。"""
    for kw, lvl in CHUNK_MARKERS:
        if kw in piece:
            return max(int(doc_clearance or 0), lvl)
    return None


def _overlaps(prev: str, cur: str) -> bool:
    """cur 是否带上一片尾部原文（chunk_text 的 overlap 是"上片尾 + 本片正文"）。

    必须按实际文本判断而不是看配置：overlap 只在 len(c) < size 时追加，
    短尾片、size 截断片都可能没有重叠，此时不该继承上一片的密级。
    """
    limit = min(len(prev), len(cur))
    return any(prev.endswith(cur[:k]) for k in range(limit, 0, -1))


def chunk_clearances(pieces: list[str], doc_clearance: int) -> list[int | None]:
    """逐片计算密级：本片标记为准，重叠带入的更高密级不得降级。

    重叠会把上一片尾部原文复制到本片开头，而密级标记（如「【绝密附注】」）
    通常落在上一片的中前部——只看本片文本时，含敏感原文的重叠片会被判低
    （实测：绝密附注正文落进 2 级切片，密级 2 的用户可直接检索到）。
    因此以「只升不降」的 floor 贯穿重叠链：一旦出现更高密级，后续有重叠的
    切片至少保持该密级。取值偏保守（宁可少给不可多给），无重叠则不继承。
    """
    out: list[int | None] = []
    floor: int | None = None
    for i, piece in enumerate(pieces):
        lvl = detect_chunk_clearance(piece, doc_clearance)
        if i and floor is not None and _overlaps(pieces[i - 1], piece):
            if lvl is None or lvl < floor:
                lvl = floor
        if lvl is not None:
            floor = lvl if floor is None else max(floor, lvl)
        out.append(lvl)
    return out


def parse_file(path: str | Path, ext: str = "") -> str:
    """解析文件为纯文本。

    解析失败一律抛 ParseError（中文可读），由上传端点转 400——
    绝不再返回 "[解析失败] ..." 伪文本：那串字会被当正文切片入库，
    污染向量库且用户无从察觉（P1 解析失败拒绝入库）。
    """
    path = Path(path)
    ext = (ext or path.suffix).lower().lstrip(".")
    if ext == "pdf":
        text, _quality = parse_pdf(path)
        return text
    try:
        if ext in ("docx", "doc"):
            from docx import Document as DocxDocument  # type: ignore
            d = DocxDocument(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)
        if ext in ("xlsx", "xls", "csv"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        if ext in ("md", "markdown", "txt", "text", "log"):
            return path.read_text(encoding="utf-8", errors="ignore")
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"文件解析失败：{e}") from e
    return path.read_text(encoding="utf-8", errors="ignore")


def parse_bytes(data: bytes, ext: str, filename: str = "") -> tuple[str, str]:
    """上传字节 → (正文, 解析质量标记)。

    质量标记：ok=正常 / low=可读比例低（疑似乱码）/ ocr=扫描件 OCR 来源。
    PDF 走 pdfx 加固解析（表格/多栏/OCR/质检）；其余格式解包后统一质检。
    """
    ext = ext.lower().lstrip(".")
    if ext in ("txt", "md", "markdown", "csv", "log", "text"):
        text = data.decode("utf-8", errors="ignore")
        return text, assess_quality(text)
    # 临时目录跟随 DATA_DIR：测试/多实例部署时各自隔离，
    # 不往仓库的 backend/data 里写运行时垃圾
    tmp_dir = DATA_DIR / "tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    # 防御性取 basename：filename 来自客户端，可能带路径穿越
    tmp = tmp_dir / (Path(filename).name if filename else f"upload.{ext or 'bin'}")
    tmp.write_bytes(data)
    try:
        if ext == "pdf":
            return parse_pdf(tmp)
        text = parse_file(tmp, ext)
        return text, assess_quality(text)
    finally:
        try:
            tmp.unlink()
        except Exception:
            pass


def split_long(text: str, size: int) -> list[str]:
    """先按句切，句子过长再硬切。"""
    pieces = [p for p in SENT_SPLIT.split(text) if p]
    out, buf = [], ""
    for p in pieces:
        if len(buf) + len(p) <= size:
            buf += p
        else:
            if buf:
                out.append(buf)
            while len(p) > size:
                out.append(p[:size])
                p = p[size:]
            buf = p
    if buf:
        out.append(buf)
    return out


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """段落优先、句子兜底的语义切片，保留重叠以维持上下文连续性。"""
    text = re.sub(r"[ \t]+", " ", text or "").strip()
    if not text:
        return []
    units: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        units.extend(split_long(para, size) if len(para) > size else [para])

    chunks, buf = [], ""
    for u in units:
        if len(buf) + len(u) + 1 <= size:
            buf = (buf + "\n" + u).strip()
        else:
            if buf:
                chunks.append(buf)
            buf = u
    if buf:
        chunks.append(buf)
    if not chunks:
        return []

    # 追加重叠：把上一段尾部内容带进下一段，避免语义被切断
    if overlap > 0 and len(chunks) > 1:
        merged = []
        for i, c in enumerate(chunks):
            if i == 0:
                merged.append(c)
                continue
            prev_tail = chunks[i - 1][-overlap:]
            merged.append(prev_tail + c if len(c) < size else c)
        chunks = merged
    return [c for c in chunks if c.strip()]


def content_hash(text: str) -> str:
    """正文指纹。版本比对用：内容没变就不该产生新版本流水。"""
    return hashlib.sha256((text or "").encode("utf-8", errors="ignore")).hexdigest()


def _last_version(db: Session, doc_id: str) -> DocumentVersion | None:
    return db.execute(
        select(DocumentVersion)
        .where(DocumentVersion.doc_id == doc_id)
        .order_by(DocumentVersion.version.desc())
        .limit(1)
    ).scalars().first()


def snapshot_version(db: Session, doc: Document, *, content: str,
                     chunk_count: int, actor: str = "") -> DocumentVersion:
    """为文档当前内容落一条版本快照（只存元数据 + 指纹，不存全文）。"""
    last = _last_version(db, doc.id)
    v = DocumentVersion(
        doc_id=doc.id,
        version=(last.version + 1) if last else 1,
        content_hash=content_hash(content),
        chunk_count=chunk_count,
        clearance=int(doc.clearance or 1),
        created_by=(actor or "")[:64],
    )
    db.add(v)
    # 版本号同步回文档主记录：引用里带的是"文档当前是第几版"，
    # 不必每次都去版本表取 max(version)。
    doc.version = v.version
    # 显式 flush：pending 的版本行立刻对同一事务可见——
    #   1. Session 是 autoflush=False，不主动 flush 时同事务内后续的
    #      _last_version() 看不到这条 pending 行，会把下一版本号算成
    #      同一个值，撞破 (doc_id, version) 唯一约束；
    #   2. 约束若被撞破，在快照点就带着堆栈暴露，而不是拖到 commit 时
    #      变成一个来源不明的 IntegrityError（那时调用栈早已出不去）。
    db.flush()
    return v


def ingest_document(
    db: Session,
    *,
    title: str,
    content: str,
    doc_type: str = "txt",
    source_name: str = "",
    owner_dept_id: str | None = None,
    clearance: int = 1,
    acl_mode: str = "dept",
    tags: str = "",
    created_by: str = "",
    parse_quality: str = "ok",
) -> Document:
    """整篇入库：切片 → 敏感信息统计 → 批量向量化 → 事务落库。

    parse_quality: 解析质量标记（ok/low/ocr），由解析链路产出；
    手动录入的文本无解析环节，默认 "ok"。
    """
    content = (content or "").strip()
    if not content:
        raise ValueError("文档内容为空")
    # 输入质量卡口（#9，单一 choke point）：schema 已拦 API 路径的超长正文，
    # 这里是内部路径（重新入档的表单 content 等）的最后一道。超限不是
    # "截断了事"——半篇文档入库比拒绝入库危险得多。
    if len(content) > MAX_DOC_CHARS:
        raise ValueError(f"文档正文超过上限（{MAX_DOC_CHARS} 字符），请拆分后上传")
    # 标题超库列宽：明确拒绝，不静默截断（调用方无从得知曾被截断）。
    if len(title) > 255:
        raise ValueError("标题过长（不得超过 255 字符）")

    doc = Document(
        title=title[:255],
        source_name=source_name[:255],
        doc_type=doc_type[:32],
        owner_dept_id=owner_dept_id,
        clearance=clearance,
        acl_mode=acl_mode if acl_mode in ("open", "dept", "restrict") else "dept",
        tags=tags[:255],
        summary=content[:300],
        pii_hits=len(scan(content)),
        # P1-4：入库侧注入扫描。命中只标记（上传端点记 decision=flag），
        # 不拦截入库——文档本身可能是"含攻击样本的演练材料"，有价值留存。
        inject_risk=scan_injection(content),
        created_by=created_by[:64],
        status="ready",
        # 白名单收敛：脏值（如空串）回落 "ok"，保证列语义只有三种取值
        parse_quality=parse_quality if parse_quality in ("ok", "low", "ocr") else "ok",
    )
    db.add(doc)
    db.flush()

    pieces = chunk_text(content)
    if not pieces:
        db.rollback()
        raise ValueError("切片结果为空")

    vecs = embedder.embed(pieces)
    levels = chunk_clearances(pieces, clearance)
    for i, (piece, vec, lvl) in enumerate(zip(pieces, vecs, levels)):
        db.add(
            Chunk(
                doc_id=doc.id, seq=i, content=piece,
                content_len=len(piece), embedding=vec,
                clearance=lvl,
            )
        )
    # P2-10：v1 版本快照。此后每次重新入档追加一条，形成可追溯的版本链。
    snapshot_version(db, doc, content=content, chunk_count=len(pieces),
                     actor=created_by)
    db.commit()
    db.refresh(doc)
    return doc


def reingest_document(
    db: Session,
    doc: Document,
    *,
    content: str,
    doc_type: str = "",
    parse_quality: str = "ok",
    actor: str = "",
) -> tuple[Document, DocumentVersion, bool]:
    """以新内容替换文档正文，并追加一条版本快照。

    返回 (doc, 版本记录, 是否真的创建了新版本)。内容指纹与当前版本一致时
    直接返回、不造版本流水——重复上传同一文件不该刷出十几个"版本"。

    旧切片直接删除：它们属于旧版本内容，留着会让检索命中已被替换掉的段落，
    等于给用户一份"自己都看不出是过期的"答案。暴露账本不动——残留条目只会
    让防拼图覆盖率偏高（更保守），删掉反而给攻击者一条刷新额度的捷径。
    """
    content = (content or "").strip()
    if not content:
        raise ValueError("文档内容为空")
    # 与 ingest_document 同一把卡口：重新入档的表单 content 不走 pydantic
    # schema，只能在这里拦（路由层 ValueError → 400 的既有转换不变）。
    if len(content) > MAX_DOC_CHARS:
        raise ValueError(f"文档正文超过上限（{MAX_DOC_CHARS} 字符），请拆分后上传")

    last = _last_version(db, doc.id)
    new_hash = content_hash(content)
    if last is not None and last.content_hash == new_hash:
        return doc, last, False

    pieces = chunk_text(content)
    if not pieces:
        raise ValueError("切片结果为空")
    vecs = embedder.embed(pieces)

    db.execute(sa_delete(Chunk).where(Chunk.doc_id == doc.id))
    levels = chunk_clearances(pieces, int(doc.clearance or 1))
    for i, (piece, vec, lvl) in enumerate(zip(pieces, vecs, levels)):
        db.add(
            Chunk(
                doc_id=doc.id, seq=i, content=piece,
                content_len=len(piece), embedding=vec,
                clearance=lvl,
            )
        )
    doc.summary = content[:300]
    doc.pii_hits = len(scan(content))
    doc.inject_risk = scan_injection(content)
    if doc_type:
        doc.doc_type = doc_type[:32]
    doc.parse_quality = (
        parse_quality if parse_quality in ("ok", "low", "ocr") else "ok"
    )

    v = snapshot_version(db, doc, content=content, chunk_count=len(pieces),
                         actor=actor)
    db.commit()
    db.refresh(doc)
    return doc, v, True
